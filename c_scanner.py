import streamlit as st
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from io import BytesIO
from st_img_pastebutton import paste
import base64
import json
import re

# ---------------------------------------------------------
# 페이지 설정
# ---------------------------------------------------------
st.set_page_config(page_title="코드 사진 & 실행결과 분석기", page_icon="💻")

# ---------------------------------------------------------
# PWA(홈 화면 앱) 설정: 앱 이름 "코드분석앱", 아이콘, 테마 색상 지정
# 정적 파일은 /static 폴더 + .streamlit/config.toml의
# enableStaticServing = true 설정이 있어야 정상 동작합니다.
# ---------------------------------------------------------
st.markdown(
    """
    <link rel="manifest" href="app/static/manifest.json">
    <meta name="theme-color" content="#2563EB">
    <link rel="apple-touch-icon" href="app/static/apple-touch-icon.png">
    <link rel="icon" href="app/static/favicon-32.png" sizes="32x32">
    <meta name="mobile-web-app-capable" content="yes">
    <meta name="apple-mobile-web-app-capable" content="yes">
    <meta name="apple-mobile-web-app-status-bar-style" content="black-translucent">
    <meta name="apple-mobile-web-app-title" content="코드분석앱">
    """,
    unsafe_allow_html=True,
)

st.title("💻 코드 사진 ➔ 코드 해설 & PPT 변환기")
st.write(
    "프로그래밍 책이나 코드가 담긴 페이지를 여러 장 찍어 올리면, AI가 어떤 언어인지 "
    "자동으로 인식해서 코드 분석과 실행 결과 풀이를 담은 여러 슬라이드의 학습 PPT를 만들어 드립니다!"
)

# ---------------------------------------------------------
# API 키 입력
# ---------------------------------------------------------
api_key = st.text_input("Gemini API 키를 입력하세요", type="password")

# ---------------------------------------------------------
# 파일 업로드 (여러 장 지원)
# ---------------------------------------------------------
uploaded_files = st.file_uploader(
    "코드가 담긴 책/자료 페이지 사진을 업로드하거나 촬영하세요 (여러 장 선택 가능)",
    type=["jpg", "jpeg", "png"],
    accept_multiple_files=True,
)

# ---------------------------------------------------------
# 클립보드에서 스크린샷 붙여넣기
# ---------------------------------------------------------
if "paste_count" not in st.session_state:
    st.session_state.paste_count = 0
if "pasted_images" not in st.session_state:
    st.session_state.pasted_images = []  # [(bytes, mime_type), ...]

st.markdown("**또는** 스크린샷을 캡처해 클립보드에 복사한 뒤, 아래 버튼을 클릭하세요")
st.caption(
    "💡 Ctrl+V가 아니라 **버튼 클릭**으로 동작합니다. "
    "Windows는 Win+Shift+S, Mac은 Cmd+Shift+4 누른 뒤 Control까지 같이 누르면 "
    "클립보드로 복사돼요. 복사 직후 바로 아래 버튼을 눌러주세요."
)
paste_col1, paste_col2 = st.columns([3, 1])

with paste_col1:
    pasted_data_url = paste(
        label="📋 클립보드 이미지 불러오기",
        key=f"paste_{st.session_state.paste_count}",
    )

with paste_col2:
    if pasted_data_url is not None:
        if st.button("➕ 이 이미지 추가"):
            header, encoded = pasted_data_url.split(",", 1)
            mime_type = header.split(";")[0].split(":")[1]
            binary_data = base64.b64decode(encoded)
            st.session_state.pasted_images.append((binary_data, mime_type))
            st.session_state.paste_count += 1  # 위젯을 새로 만들어 다음 붙여넣기를 받음
            st.rerun()

if st.session_state.pasted_images:
    if st.button("🗑️ 붙여넣은 이미지 전체 삭제"):
        st.session_state.pasted_images = []
        st.rerun()

# ---------------------------------------------------------
# 업로드 파일 + 붙여넣은 이미지 미리보기
# ---------------------------------------------------------
preview_items = []  # [(bytes, mime, caption), ...]
if uploaded_files:
    preview_items += [(f.getvalue(), f.type, f"업로드 {i + 1}") for i, f in enumerate(uploaded_files)]
preview_items += [
    (data, mime, f"붙여넣기 {i + 1}") for i, (data, mime) in enumerate(st.session_state.pasted_images)
]

if preview_items:
    cols = st.columns(min(len(preview_items), 4))
    for i, (data, _, caption) in enumerate(preview_items):
        with cols[i % len(cols)]:
            st.image(data, caption=caption, use_container_width=True)


# ---------------------------------------------------------
# Gemini에게 "구조화된 JSON"으로 응답하도록 요청
# -> 슬라이드를 여러 장으로 나눠 만들기 위함
# ---------------------------------------------------------
def analyze_code_pages(client: genai.Client, images: list[bytes], mime_types: list[str]) -> dict:
    prompt = (
        "당신은 프로그래밍 책의 예제 코드를 학습 슬라이드로 정리해주는 도우미입니다. "
        "첨부된 책/자료 페이지 사진들을 분석해서, 코딩 초보자가 영어 문장을 "
        "단어별로 해석하듯 코드를 한 줄씩 따라가며 이해할 수 있는 학습 노트를 만들어주세요. "
        "코드가 어떤 프로그래밍 언어로 작성되었는지는 사진을 보고 스스로 판단하세요 "
        "(C, C++, Java, Python, JavaScript 등 어떤 언어든 가능합니다).\n\n"
        "다음 JSON 형식으로만 응답하세요. 다른 설명, 마크다운 코드블록(```) "
        "표시는 절대 포함하지 마세요:\n\n"
        "{\n"
        '  "title": "전체 학습 노트 제목 (예: OO 예제 코드 분석 노트)",\n'
        '  "language": "코드의 프로그래밍 언어를 나타내는 짧은 식별자 하나. 다음 중에서 "\n'
        '              "고르세요: python, c, cpp, java, javascript, typescript, csharp, "\n'
        '              "go, rust, kotlin, swift, php, ruby, sql, html, text. 확신이 서지 "\n'
        '              "않으면 text로 표기하세요.",\n'
        '  "libraries": [\n'
        "    {\n"
        '      "name": "코드에 포함된 import/include/require 등 라이브러리·모듈 이름",\n'
        '      "description": "이 라이브러리가 왜 필요한지, 어떤 기능을 제공하는지 한 문장으로 설명"\n'
        "    }\n"
        "  ],\n"
        '  "sections": [\n'
        "    {\n"
        '      "heading": "슬라이드 제목",\n'
        '      "bullets": ["요점 1", "요점 2", "요점 3"]\n'
        "    }\n"
        "  ],\n"
        '  "code_lines": [\n'
        "    {\n"
        '      "code": "사진 속 코드 원문 한 줄 또는 한 구문",\n'
        '      "explanation": "이 코드 한 줄이 정확히 무엇을 하는지 우리말로 해설"\n'
        "    }\n"
        "  ]\n"
        "}\n\n"
        "각 항목별 작성 규칙:\n"
        "- language: 코드 문법(세미콜론, 들여쓰기 방식, 키워드 등)을 보고 가장 가능성 높은 언어 "
        "하나만 고르세요.\n"
        "- libraries: 코드 상단에 import/#include/require/using 등으로 쓰인 라이브러리나 모듈을 "
        "모두 나열하세요. 없으면 빈 배열([])로 두세요.\n"
        "- sections: 반드시 다음 3가지로 구성하세요.\n"
        "  1. 코드 핵심 목적 및 문법 설명\n"
        "  2. 주요 함수(또는 메서드) 및 연산자 풀이\n"
        "  3. 예상 실행 결과 및 화면 설명\n"
        "  각 섹션의 bullets는 2~5개, 간결한 문장으로 작성하세요.\n"
        "- code_lines: 사진 속 코드에서 의미 있는 구문(변수 선언, 함수/메서드 호출, 조건문, "
        "반복문, 연산 등)을 코드에 나온 순서 그대로, 최대한 빠짐없이 나열하세요. 중괄호나 "
        "들여쓰기만 있는 줄, 빈 줄은 생략해도 되지만, 그 외에는 최대한 모든 실행 라인을 "
        "포함하세요 (보통 8~20줄). code 필드는 사진에 보이는 코드 그대로(들여쓰기, 구두점 포함) "
        "적고, explanation은 그 줄이 프로그램에서 실제로 하는 동작을 초보자도 이해할 수 있게 "
        "풀어서 설명하세요.\n\n"
        "특수문자(**, ###, - 등 마크다운 기호)는 bullets와 explanation에 쓰지 마세요. "
        "변수명, 함수명, 코드 자체는 원문 그대로 써도 됩니다."
    )

    parts = [
        types.Part.from_bytes(data=img, mime_type=mt)
        for img, mt in zip(images, mime_types)
    ]
    parts.append(prompt)

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=parts,
    )

    raw_text = (response.text or "").strip()
    if not raw_text:
        raise ValueError(
            "AI가 빈 응답을 반환했습니다. 안전 필터에 걸렸거나 이미지를 "
            "인식하지 못했을 수 있습니다. 다른 사진으로 다시 시도해보세요."
        )

    # 혹시 모델이 ```json ... ``` 코드블록으로 감싸서 보낼 경우 제거
    cleaned = re.sub(r"^```(?:json)?|```$", "", raw_text.strip(), flags=re.MULTILINE).strip()

    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        # JSON 파싱 실패 시, 최소한 하나의 슬라이드로라도 보여주기 위한 대비책
        data = {
            "title": "코드 학습 및 분석 노트",
            "language": "text",
            "libraries": [],
            "sections": [{"heading": "분석 결과", "bullets": [raw_text]}],
            "code_lines": [],
        }

    return data


# ---------------------------------------------------------
# 텍스트 프레임에 자동 줄맞춤 + 글자 수에 따른 폰트 크기 조절
# ---------------------------------------------------------
def set_autofit_text(text_frame, bullets: list[str]):
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # 내용이 적어도 위에 쏠리지 않고 중앙에 배치

    total_chars = sum(len(b) for b in bullets)

    # 글자 수가 적을수록(= 슬라이드가 휑해 보일수록) 글자를 크게, 줄 간격도 넓게
    if total_chars > 500:
        font_size, space_after = Pt(14), Pt(8)
    elif total_chars > 250:
        font_size, space_after = Pt(16), Pt(12)
    elif total_chars > 100:
        font_size, space_after = Pt(18), Pt(18)
    else:
        font_size, space_after = Pt(22), Pt(24)

    text_frame.clear()
    for i, bullet in enumerate(bullets):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = font_size
        p.space_after = space_after


# ---------------------------------------------------------
# "라이브러리 설명" 슬라이드 추가
# ---------------------------------------------------------
def add_library_slide(prs, libraries: list[dict]):
    if not libraries:
        return
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "사용된 라이브러리/모듈 설명"
    for p in slide.shapes.title.text_frame.paragraphs:
        p.font.size = Pt(24)
        p.font.bold = True

    body = slide.placeholders[1]
    bullets = [
        f"{lib.get('name', '이름 없음')} : {lib.get('description', '')}"
        for lib in libraries
    ]
    set_autofit_text(body.text_frame, bullets)


# ---------------------------------------------------------
# 항목을 슬라이드 수에 맞춰 "균등하게" 나눠줌
# (마지막 슬라이드에 1~2개만 남아 휑해지는 것을 방지)
# ---------------------------------------------------------
def chunk_evenly(items: list, max_per_slide: int) -> list[list]:
    n = len(items)
    if n == 0:
        return []
    num_slides = -(-n // max_per_slide)  # 올림 나눗셈
    per_slide = -(-n // num_slides)      # 슬라이드 수에 맞게 다시 균등 배분
    return [items[i:i + per_slide] for i in range(0, n, per_slide)]


# ---------------------------------------------------------
# "코드 한 줄씩 해설" 슬라이드 추가 (코드 → 설명 쌍을 함께 표시)
# ---------------------------------------------------------
def add_code_breakdown_slides(prs, code_lines: list[dict], max_lines_per_slide: int = 6):
    if not code_lines:
        return
    layout = prs.slide_layouts[1]
    chunks = chunk_evenly(code_lines, max_lines_per_slide)

    line_no = 1
    for chunk in chunks:
        slide = prs.slides.add_slide(layout)
        end = line_no + len(chunk) - 1
        slide.shapes.title.text = f"코드 한 줄씩 해설 ({line_no}~{end}번째)"
        for p in slide.shapes.title.text_frame.paragraphs:
            p.font.size = Pt(22)
            p.font.bold = True
        line_no = end + 1

        tf = slide.placeholders[1].text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE  # 줄 수가 적어도 중앙에 배치
        tf.clear()

        # 줄 수가 적을수록 글자를 키우고 줄 간격을 넓혀서 빈 공간을 채움
        if len(chunk) <= 2:
            font_size, space_after = Pt(18), Pt(28)
        elif len(chunk) <= 4:
            font_size, space_after = Pt(15), Pt(20)
        else:
            font_size, space_after = Pt(13), Pt(14)

        for i, item in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = space_after

            code_run = p.add_run()
            code_run.text = item.get("code", "")
            code_run.font.name = "Consolas"
            code_run.font.size = font_size
            code_run.font.bold = True
            code_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)

            arrow_run = p.add_run()
            arrow_run.text = "  →  "
            arrow_run.font.size = font_size

            expl_run = p.add_run()
            expl_run.text = item.get("explanation", "")
            expl_run.font.size = font_size


# ---------------------------------------------------------
# python-pptx로 여러 슬라이드 PPT 생성 (메모리에서 바로 생성)
# ---------------------------------------------------------
def build_pptx(data: dict) -> BytesIO:
    prs = Presentation()

    # 1. 표지 슬라이드
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = data.get("title", "코드 학습 및 분석 노트")
    for p in slide.shapes.title.text_frame.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "AI가 자동으로 생성한 코드 학습 노트"

    # 2. 라이브러리 설명 슬라이드
    add_library_slide(prs, data.get("libraries", []))

    # 3. 섹션별 내용 슬라이드 (핵심 목적 / 함수 풀이 / 실행 결과)
    content_layout = prs.slide_layouts[1]
    for section in data.get("sections", []):
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = section.get("heading", "내용")
        for p in slide.shapes.title.text_frame.paragraphs:
            p.font.size = Pt(24)
            p.font.bold = True

        body = slide.placeholders[1]
        bullets = section.get("bullets", []) or ["내용 없음"]
        set_autofit_text(body.text_frame, bullets)

    # 4. 코드 한 줄씩 해설 슬라이드 (코드 → 설명, 영어 직독직해처럼 매칭)
    add_code_breakdown_slides(prs, data.get("code_lines", []))

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------
# 메인 로직
# ---------------------------------------------------------
if preview_items and api_key:
    if st.button("코드 분석 및 PPT 변환 시작"):
        with st.spinner("AI가 코드를 해설하고 실행 결과를 분석하는 중입니다..."):
            try:
                client = genai.Client(api_key=api_key)

                images = [item[0] for item in preview_items]
                mime_types = [item[1] for item in preview_items]

                data = analyze_code_pages(client, images, mime_types)
                pptx_buffer = build_pptx(data)

                st.success("코드 분석 및 PPT 생성 완료!")

                st.download_button(
                    label="📥 코드 학습 PPT 파일 다운로드",
                    data=pptx_buffer,
                    file_name="code_study_note.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )

                st.markdown("---")
                st.markdown(f"### 📝 {data.get('title', '코드 해설 및 실행 결과 풀이')}")

                libraries = data.get("libraries", [])
                if libraries:
                    st.markdown("#### 📚 사용된 라이브러리/모듈")
                    for lib in libraries:
                        st.markdown(f"- **{lib.get('name', '')}** : {lib.get('description', '')}")

                for section in data.get("sections", []):
                    st.markdown(f"**{section.get('heading', '')}**")
                    for b in section.get("bullets", []):
                        st.markdown(f"- {b}")

                code_lines = data.get("code_lines", [])
                if code_lines:
                    detected_language = data.get("language", "text") or "text"
                    st.markdown(f"#### 🔍 코드 한 줄씩 해설 (감지된 언어: {detected_language})")
                    for item in code_lines:
                        code_col, expl_col = st.columns([1, 1])
                        with code_col:
                            st.code(item.get("code", ""), language=detected_language)
                        with expl_col:
                            st.markdown(item.get("explanation", ""))

            except Exception as e:
                st.error(f"오류가 발생했습니다: {e}")

elif preview_items and not api_key:
    st.warning("위쪽 빈칸에 Gemini API 키를 먼저 입력해 주세요!")
elif api_key and not preview_items:
    st.info("코드가 담긴 사진을 한 장 이상 업로드하거나 붙여넣어 주세요.")
