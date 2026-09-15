import streamlit as st
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Pt
from io import BytesIO
from st_img_pastebutton import paste
import base64
import json
import re

# ---------------------------------------------------------
# 페이지 설정
# ---------------------------------------------------------
st.set_page_config(page_title="책 사진 PPT 변환기", page_icon="📚")

st.title("📚 책 사진 ➔ 내용 요약 및 파워포인트(PPT) 변환기")
st.write(
    "책 페이지를 사진으로 여러 장 올리면, AI가 내용을 구조화해서 "
    "여러 슬라이드로 구성된 PPT 파일을 만들어드립니다."
)

# ---------------------------------------------------------
# API 키 입력
# ---------------------------------------------------------
api_key = st.text_input("Gemini API 키를 입력하세요", type="password")

# ---------------------------------------------------------
# 파일 업로드 (여러 장 지원)
# ---------------------------------------------------------
uploaded_files = st.file_uploader(
    "책 사진을 업로드하거나 촬영하세요 (여러 장 선택 가능)",
    type=["jpg", "jpeg", "png"],
    accept_multiple_files=True,
)

# ---------------------------------------------------------
# 클립보드에서 스크린샷 붙여넣기
# ---------------------------------------------------------
if "paste_count" not in st.session_state:
    st.session_state.paste_count = 0
if "pasted_images" not in st.session_state:
    st.session_state.pasted_images = []  # [(bytes, mime_type), ...]

st.markdown("**또는** 스크린샷을 캡처해 클립보드에 복사한 뒤, 아래 버튼을 클릭하세요")
st.caption(
    "💡 Ctrl+V가 아니라 **버튼 클릭**으로 동작합니다. "
    "Windows는 Win+Shift+S, Mac은 Cmd+Shift+4 누른 뒤 Control까지 같이 누르면 "
    "클립보드로 복사돼요. 복사 직후 바로 아래 버튼을 눌러주세요."
)
paste_col1, paste_col2 = st.columns([3, 1])

with paste_col1:
    pasted_data_url = paste(
        label="📋 클립보드 이미지 불러오기",
        key=f"paste_{st.session_state.paste_count}",
    )

with paste_col2:
    if pasted_data_url is not None:
        if st.button("➕ 이 이미지 추가"):
            header, encoded = pasted_data_url.split(",", 1)
            mime_type = header.split(";")[0].split(":")[1]
            binary_data = base64.b64decode(encoded)
            st.session_state.pasted_images.append((binary_data, mime_type))
            st.session_state.paste_count += 1  # 위젯을 새로 만들어 다음 붙여넣기를 받음
            st.rerun()

if st.session_state.pasted_images:
    if st.button("🗑️ 붙여넣은 이미지 전체 삭제"):
        st.session_state.pasted_images = []
        st.rerun()

# ---------------------------------------------------------
# 업로드 파일 + 붙여넣은 이미지 미리보기
# ---------------------------------------------------------
preview_items = []  # [(bytes, mime, caption), ...]
if uploaded_files:
    preview_items += [(f.getvalue(), f.type, f"업로드 {i + 1}") for i, f in enumerate(uploaded_files)]
preview_items += [
    (data, mime, f"붙여넣기 {i + 1}") for i, (data, mime) in enumerate(st.session_state.pasted_images)
]

if preview_items:
    cols = st.columns(min(len(preview_items), 4))
    for i, (data, _, caption) in enumerate(preview_items):
        with cols[i % len(cols)]:
            st.image(data, caption=caption, use_container_width=True)


# ---------------------------------------------------------
# Gemini에게 "구조화된 JSON"으로 응답하도록 요청
# -> 슬라이드를 여러 장으로 나눠 만들기 위함
# ---------------------------------------------------------
def analyze_book_pages(client: genai.Client, images: list[bytes], mime_types: list[str]) -> dict:
    prompt = (
        "당신은 책 내용을 파워포인트 슬라이드로 정리해주는 도우미입니다. "
        "첨부된 책 페이지 사진들을 분석해서 학습 노트를 만들어주세요.\n\n"
        "다음 JSON 형식으로만 응답하세요. 다른 설명, 마크다운 코드블록(```) "
        "표시는 절대 포함하지 마세요:\n\n"
        "{\n"
        '  "title": "전체 학습 노트 제목 (짧게)",\n'
        '  "sections": [\n'
        "    {\n"
        '      "heading": "슬라이드 제목 (예: 핵심 주제 및 개요)",\n'
        '      "bullets": ["요점 1", "요점 2", "요점 3"]\n'
        "    }\n"
        "  ]\n"
        "}\n\n"
        "섹션은 보통 3~5개 정도로 구성하세요 "
        "(예: 핵심 주제 및 개요 / 주요 내용 및 개념 / 핵심 요약 및 시사점 등). "
        "각 섹션의 bullets는 2~5개, 한 줄에 너무 길지 않게 간결한 문장으로 작성하세요. "
        "특수문자(**, ###, - 등)는 쓰지 마세요."
    )

    parts = [
        types.Part.from_bytes(data=img, mime_type=mt)
        for img, mt in zip(images, mime_types)
    ]
    parts.append(prompt)

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=parts,
    )

    raw_text = (response.text or "").strip()
    if not raw_text:
        raise ValueError(
            "AI가 빈 응답을 반환했습니다. 안전 필터에 걸렸거나 이미지를 "
            "인식하지 못했을 수 있습니다. 다른 사진으로 다시 시도해보세요."
        )

    # 혹시 모델이 ```json ... ``` 코드블록으로 감싸서 보낼 경우 제거
    cleaned = re.sub(r"^```(?:json)?|```$", "", raw_text.strip(), flags=re.MULTILINE).strip()

    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        # JSON 파싱 실패 시, 최소한 하나의 슬라이드로라도 보여주기 위한 대비책
        data = {
            "title": "AI 학습 요약 노트",
            "sections": [{"heading": "요약", "bullets": [raw_text]}],
        }

    return data


# ---------------------------------------------------------
# 텍스트 프레임에 자동 줄맞춤 + 글자 수에 따른 폰트 크기 조절
# ---------------------------------------------------------
def set_autofit_text(text_frame, bullets: list[str]):
    text_frame.word_wrap = True
    total_chars = sum(len(b) for b in bullets)

    if total_chars > 500:
        font_size = Pt(14)
    elif total_chars > 250:
        font_size = Pt(16)
    else:
        font_size = Pt(18)

    text_frame.clear()
    for i, bullet in enumerate(bullets):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = font_size


# ---------------------------------------------------------
# python-pptx로 여러 슬라이드 PPT 생성 (메모리에서 바로 생성)
# ---------------------------------------------------------
def build_pptx(data: dict) -> BytesIO:
    prs = Presentation()

    # 1. 표지 슬라이드
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = data.get("title", "AI 학습 요약 노트")
    for p in slide.shapes.title.text_frame.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "AI가 자동으로 생성한 학습 노트"

    # 2. 섹션별 내용 슬라이드
    content_layout = prs.slide_layouts[1]
    for section in data.get("sections", []):
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = section.get("heading", "내용")
        for p in slide.shapes.title.text_frame.paragraphs:
            p.font.size = Pt(24)
            p.font.bold = True

        body = slide.placeholders[1]
        bullets = section.get("bullets", []) or ["내용 없음"]
        set_autofit_text(body.text_frame, bullets)

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------
# 메인 로직
# ---------------------------------------------------------
if preview_items and api_key:
    if st.button("AI 분석 및 PPT 변환 시작"):
        with st.spinner("AI가 책을 분석하고 PPT를 만드는 중입니다..."):
            try:
                client = genai.Client(api_key=api_key)

                images = [item[0] for item in preview_items]
                mime_types = [item[1] for item in preview_items]

                data = analyze_book_pages(client, images, mime_types)
                pptx_buffer = build_pptx(data)

                st.success("분석 완료!")

                st.download_button(
                    label="📥 세련된 요약 PPT 파일 다운로드",
                    data=pptx_buffer,
                    file_name="book_study_note.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )

                st.markdown("---")
                st.markdown(f"### 📝 {data.get('title', 'AI 요약 결과')}")
                for section in data.get("sections", []):
                    st.markdown(f"**{section.get('heading', '')}**")
                    for b in section.get("bullets", []):
                        st.markdown(f"- {b}")

            except Exception as e:
                st.error(f"오류가 발생했습니다: {e}")

elif preview_items and not api_key:
    st.warning("위쪽 빈칸에 Gemini API 키를 먼저 입력해 주세요!")
elif api_key and not preview_items:
    st.info("책 사진을 한 장 이상 업로드하거나 붙여넣어 주세요.")
